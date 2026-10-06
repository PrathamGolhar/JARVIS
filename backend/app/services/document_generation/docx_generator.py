from datetime import datetime
from pathlib import Path
import re
from uuid import uuid4

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.settings import settings


def generate_docx_document(
    title: str,
    content_markdown: str,
    subtitle: str = "",
    author: str = "JARVIS AI Core",
    filename: str | None = None,
) -> Path:
    """
    Generate a formatted Microsoft Word (.docx) document.
    """
    out_dir = settings.generated_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_title = re.sub(r"[^\w\-_]", "_", title)[:40]
    out_filename = filename or f"{safe_title}_{uuid4().hex[:6]}.docx"
    dest_path = out_dir / out_filename

    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = f"JARVIS Intelligence Report | {title}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(148, 163, 184)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = f"Confidential & Prepared for Master — {datetime.now().strftime('%Y-%m-%d')}"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(148, 163, 184)

    # Document Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(4)

    if subtitle:
        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run(subtitle)
        r_sub.italic = True
        r_sub.font.size = Pt(13)
        r_sub.font.color.rgb = RGBColor(2, 132, 199)
        p_sub.paragraph_format.space_after = Pt(12)

    # Metadata Line
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Author: {author}  •  Date: {datetime.now().strftime('%B %d, %Y')}")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.paragraph_format.space_after = Pt(18)

    # Markdown Parsing
    lines = content_markdown.splitlines()
    in_code = False
    code_lines = []

    for line in lines:
        s = line.strip()

        if s.startswith("```"):
            if in_code:
                # Add code block
                table = doc.add_table(rows=1, cols=1)
                table.autofit = True
                cell = table.cell(0, 0)
                shading_elm = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading_elm)
                cp = cell.paragraphs[0]
                c_run = cp.add_run("\n".join(code_lines))
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(9)
                c_run.font.color.rgb = RGBColor(15, 23, 42)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not s:
            continue

        if s.startswith("# "):
            h = doc.add_heading(s[2:], level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
        elif s.startswith("## "):
            h = doc.add_heading(s[3:], level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
        elif s.startswith("### "):
            h = doc.add_heading(s[4:], level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
        elif s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(s[2:], style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(re.sub(r"^\d+\.\s", "", s), style="List Number")
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            # Clean markdown bold/italics
            text = s
            parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("*") and part.endswith("*"):
                    r = p.add_run(part[1:-1])
                    r.italic = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.color.rgb = RGBColor(2, 132, 199)
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(6)

    doc.save(str(dest_path))
    return dest_path

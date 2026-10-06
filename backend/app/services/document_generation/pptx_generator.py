from pathlib import Path
import re
from uuid import uuid4

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas import SlideContent
from app.settings import settings

THEME_PALETTES = {
    "cyan": {
        "bg": RGBColor(10, 25, 47),
        "primary": RGBColor(100, 255, 218),
        "text": RGBColor(204, 214, 246),
        "card": RGBColor(17, 34, 64),
    },
    "gold": {
        "bg": RGBColor(24, 24, 27),
        "primary": RGBColor(250, 204, 21),
        "text": RGBColor(244, 244, 245),
        "card": RGBColor(39, 39, 42),
    },
    "emerald": {
        "bg": RGBColor(6, 78, 59),
        "primary": RGBColor(52, 211, 153),
        "text": RGBColor(236, 253, 245),
        "card": RGBColor(4, 120, 87),
    },
    "crimson": {
        "bg": RGBColor(69, 10, 10),
        "primary": RGBColor(248, 113, 113),
        "text": RGBColor(254, 242, 242),
        "card": RGBColor(127, 29, 29),
    },
    "purple": {
        "bg": RGBColor(46, 16, 101),
        "primary": RGBColor(192, 132, 252),
        "text": RGBColor(250, 245, 255),
        "card": RGBColor(88, 28, 135),
    },
}


def generate_pptx_presentation(
    topic: str,
    slides: list[SlideContent],
    theme_color: str = "cyan",
    filename: str | None = None,
) -> Path:
    """
    Generate a modern 16:9 widescreen PowerPoint presentation.
    """
    out_dir = settings.generated_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_topic = re.sub(r"[^\w\-_]", "_", topic)[:40]
    out_filename = filename or f"{safe_topic}_{uuid4().hex[:6]}.pptx"
    dest_path = out_dir / out_filename

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    palette = THEME_PALETTES.get(theme_color.lower(), THEME_PALETTES["cyan"])
    blank_layout = prs.slide_layouts[6]

    # 1. Title Slide
    title_slide = prs.slides.add_slide(blank_layout)
    bg_shape = title_slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = palette["bg"]
    bg_shape.line.fill.background()

    # Title box
    tx_box = title_slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = topic
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = palette["primary"]
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.LEFT

    p2 = tf.add_paragraph()
    p2.text = "Prepared by JARVIS Artificial Intelligence Core"
    p2.font.size = Pt(20)
    p2.font.color.rgb = palette["text"]
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(16)

    # 2. Content Slides
    for slide_data in slides:
        slide = prs.slides.add_slide(blank_layout)

        # Background
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = palette["bg"]
        bg.line.fill.background()

        # Header bar
        header_card = slide.shapes.add_shape(1, Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.1))
        header_card.fill.solid()
        header_card.fill.fore_color.rgb = palette["card"]
        header_card.line.fill.background()

        # Slide Title
        tbox = slide.shapes.add_textbox(Inches(1.1), Inches(0.7), Inches(11.0), Inches(0.9))
        tt_frame = tbox.text_frame
        tt_frame.word_wrap = True
        tp = tt_frame.paragraphs[0]
        tp.text = slide_data.title
        tp.font.bold = True
        tp.font.size = Pt(28)
        tp.font.color.rgb = palette["primary"]
        tp.font.name = "Segoe UI"

        # Content Card
        content_card = slide.shapes.add_shape(1, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.8))
        content_card.fill.solid()
        content_card.fill.fore_color.rgb = palette["card"]
        content_card.line.fill.background()

        cbox = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(11.0), Inches(4.2))
        cf = cbox.text_frame
        cf.word_wrap = True

        if slide_data.subtitle:
            sp = cf.paragraphs[0]
            sp.text = slide_data.subtitle
            sp.font.italic = True
            sp.font.size = Pt(18)
            sp.font.color.rgb = palette["primary"]
            sp.space_after = Pt(12)

        for i, bullet in enumerate(slide_data.bullets):
            bp = cf.add_paragraph() if (slide_data.subtitle or i > 0) else cf.paragraphs[0]
            bp.text = f"•   {bullet}"
            bp.font.size = Pt(18)
            bp.font.color.rgb = palette["text"]
            bp.font.name = "Segoe UI"
            bp.space_before = Pt(8)
            bp.space_after = Pt(8)

        # Speaker Notes
        if slide_data.speaker_notes:
            notes_slide = slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = slide_data.speaker_notes

    prs.save(str(dest_path))
    return dest_path
